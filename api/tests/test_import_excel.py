"""Phase 5: customer spreadsheet import — messy input, clean report, idempotent."""

from __future__ import annotations

import csv
import random
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from api.db.models import AuditLog, Customer, Tenant
from api.db.session import Database
from api.onboarding.importer import Mapping, clean_name, name_key, plan, suggest
from api.onboarding.phones import normalize_phone
from api.onboarding.sheets import SheetError, read_table
from api.scripts import import_excel
from api.tests.conftest import TenantPair, make_tenant

# ---------------------------------------------------------------- phones


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("0501234567", "971501234567"),
        ("050 123 4567", "971501234567"),
        ("050-123-4567", "971501234567"),
        ("(050) 123.4567", "971501234567"),
        ("+971 50 123 4567", "971501234567"),
        ("00971501234567", "971501234567"),
        ("971501234567", "971501234567"),
        ("+971 050 123 4567", "971501234567"),
        ("501234567", "971501234567"),
        (501234567, "971501234567"),
        (501234567.0, "971501234567"),
        (971501234567.0, "971501234567"),
        ("5.01234567E+08", "971501234567"),
        ("٠٥٠١٢٣٤٥٦٧", "971501234567"),
        ("۰۵۵۱۲۳۴۵۶۷", "971551234567"),
        ("04 123 4567 / 055 765 4321", "971557654321"),
        ("0521112222 or 0563334444", "971521112222"),
        (" 0581234567 ", "971581234567"),
    ],
)
def test_phone_normalises_to_e164(raw: object, expected: str) -> None:
    assert normalize_phone(raw).wa_id == expected


@pytest.mark.parametrize(
    ("raw", "reason"),
    [
        (None, "no phone number"),
        ("", "no phone number"),
        ("n/a", "no phone number"),
        ("04 123 4567", "landline number (not on WhatsApp)"),
        ("+971 4 123 4567", "landline number (not on WhatsApp)"),
        ("0591234567", "invalid UAE mobile number"),
        ("+44 7700 900123", "country not supported"),
        ("12345", "unrecognised phone number"),
        (True, "no phone number"),
    ],
)
def test_phone_rejections_carry_a_reason(raw: object, reason: str) -> None:
    p = normalize_phone(raw)
    assert p.wa_id is None
    assert p.reason == reason


def test_names_and_keys() -> None:
    assert clean_name("  AHMED   ALI ") == "Ahmed Ali"
    assert clean_name("محمد  أحمد") == "محمد أحمد"
    assert clean_name("-") is None
    assert clean_name("N/A") is None
    assert clean_name("McDonald") == "McDonald"  # mixed case is left alone
    assert name_key("أحمد علي") == name_key("علي احمد")
    assert name_key("Ahmed ALI") == name_key("ali ahmed")


def test_suggested_mapping_reads_english_and_arabic_headers() -> None:
    got = suggest(["Customer Name", "Mobile No.", "رقم الجوال", "Area", "Emirate", "Villa / Flat",
                   "Cust ID", "Balance"])  # fmt: skip
    assert got == {
        "Customer Name": "name",
        "Mobile No.": "phone",
        "رقم الجوال": "phone",
        "Area": "area",
        "Emirate": "emirate",
        "Villa / Flat": "address_note",
        "Cust ID": "external_ref",
        "Balance": "ignore",
    }


# ---------------------------------------------------------------- messy workbook


def _messy_workbook(path: Path) -> None:
    wb = Workbook()
    notes = wb.active
    assert notes is not None
    notes.title = "Notes"
    notes.sheet_state = "hidden"
    notes["A1"] = "internal"
    ws = wb.create_sheet("Customers")
    ws["A1"] = "Customer list — exported 2026"
    ws.merge_cells("A1:F1")  # title merged across the table: never mistaken for the header
    ws["A2"] = None
    ws.append(["Cust ID", "Customer Name", "Mobile No.", "Area", "Emirate", "Villa / Flat"])
    rows: list[list[Any]] = [
        ["C1", "AHMED ALI", 501234567.0, "Al Nahda", "dubai", "Villa 12"],
        ["C2", "محمد أحمد", "٠٥٥١٢٣٤٥٦٧", None, "الشارقة", None],  # area from the merge below
        ["C3", "Sara Khan", "5.21112222E+08", None, "SHJ", "Flat 3"],
        [None, None, None, None, None, None],  # blank
        ["C4", "Landline Co", "04 123 4567", "Deira", "Dubai", None],
        [
            "Cust ID",
            "Customer Name",
            "Mobile No.",
            "Area",
            "Emirate",
            "Villa / Flat",
        ],  # page header
        ["C5", "Ahmed Ali", "+971 50 123 4567", "Al Nahda", "Dubai", "Villa 12, Street 4"],  # dup
        ["C6", "No Phone", None, "JLT", "Dubai", None],
        ["C7", "Omar Hassan", "0561112233", "Al Nahda", "Dubai", None],
        ["C8", "Omar Hasan", "0567778899", "Al Nahda", "Dubai", None],  # similar name, other no.
    ]
    for r in rows:
        ws.append(r)
    ws.merge_cells("D5:D6")  # "Al Qasimia" applies to rows 5 and 6
    ws["D5"] = "Al Qasimia"
    wb.save(path)


MESSY_MAPPING = Mapping(
    columns={
        "Cust ID": "external_ref",
        "Customer Name": "name",
        "Mobile No.": "phone",
        "Area": "area",
        "Emirate": "emirate",
        "Villa / Flat": "address_note",
    }
)


def test_messy_sheet_is_read_and_planned(tmp_path: Path) -> None:
    path = tmp_path / "messy.xlsx"
    _messy_workbook(path)
    table = read_table(path)
    assert table.sheet == "Customers"  # the hidden sheet is skipped
    assert table.header_row == 3
    assert table.blank_rows == 1
    assert table.repeated_headers == 1

    p = plan(table, MESSY_MAPPING)
    by_phone = {c.wa_id: c for c in p.candidates}
    assert set(by_phone) == {
        "971501234567",
        "971551234567",
        "971521112222",
        "971561112233",
        "971567778899",
    }
    ahmed = by_phone["971501234567"]
    assert ahmed.values["name"] == "Ahmed Ali"
    assert ahmed.values["emirate"] == "Dubai"
    assert ahmed.values["address_note"] == "Villa 12"  # first row wins
    assert ahmed.rows == [4, 10]
    assert any("different address note" in c for c in ahmed.conflicts)
    assert by_phone["971551234567"].values == {
        "external_ref": "C2",
        "name": "محمد أحمد",
        "area": "Al Qasimia",
        "emirate": "Sharjah",
    }
    assert by_phone["971521112222"].values["area"] == "Al Qasimia"
    assert p.duplicate_rows == 1
    assert sorted((r.row, r.reason) for r in p.rejected) == [
        (8, "landline number (not on WhatsApp)"),
        (11, "no phone number"),
    ]
    assert p.possible_duplicates == [
        ("971561112233", "971567778899", "similar names (95%) in the same area")
    ]


def test_csv_in_windows_arabic_encoding(tmp_path: Path) -> None:
    path = tmp_path / "list.csv"
    path.write_bytes("الاسم;الجوال\nمحمد;0501112222\n".encode("cp1256"))
    table = read_table(path)
    assert table.columns == ["الاسم", "الجوال"]
    p = plan(table, Mapping(columns={"الاسم": "name", "الجوال": "phone"}))
    assert p.candidates[0].values["name"] == "محمد"


def test_unreadable_inputs_are_explained(tmp_path: Path) -> None:
    with pytest.raises(SheetError, match=r"save as \.xlsx"):
        read_table(tmp_path / "old.xls")
    path = tmp_path / "x.csv"
    path.write_text("only\n1\n")
    with pytest.raises(SheetError, match="no header row"):
        read_table(path)
    path.write_text("a,b\n1,2\n")
    with pytest.raises(ValueError, match="no phone column"):
        plan(read_table(path), Mapping(columns={"a": "name"}))
    with pytest.raises(ValueError, match="does not have"):
        plan(read_table(path), Mapping(columns={"zzz": "phone"}))


# ---------------------------------------------------------------- against the database


async def _slug(db: Database, tenant_id: uuid.UUID) -> str:
    async with db.platform_session() as s:
        slug = await s.scalar(select(Tenant.slug).where(Tenant.id == tenant_id))
    assert slug is not None
    return slug


async def _customers(db: Database, tenant_id: uuid.UUID) -> dict[str, Customer]:
    async with db.tenant_session(tenant_id) as s:
        return {c.wa_id or "": c for c in (await s.scalars(select(Customer))).all()}


async def test_import_fills_blanks_keeps_corrections_and_is_idempotent(
    db: Database, tenants: TenantPair, tmp_path: Path
) -> None:
    path = tmp_path / "messy.xlsx"
    _messy_workbook(path)
    table = read_table(path)
    slug = await _slug(db, tenants.a)
    async with db.tenant_session(tenants.a) as s:
        # already known: the customer told the agent their area; the sheet has a different one
        s.add(Customer(wa_id="971501234567", name=None, area="Al Qusais", opt_in_status="opted_in"))
        s.add(Customer(wa_id="971561112233", name="Omar", opt_in_status="opted_out"))
    async with db.tenant_session(tenants.b) as s:
        s.add(Customer(wa_id="971551234567", name="B's customer"))

    # dry run: counts, no writes
    _, dry = await import_excel.run(
        db, slug, table, MESSY_MAPPING, file_sha256="0" * 64, commit=False
    )
    assert (dry.imported, dry.updated) == (3, 2)
    assert len(await _customers(db, tenants.a)) == 3  # 2 seeded + the fixture's customer

    _, out = await import_excel.run(
        db, slug, table, MESSY_MAPPING, file_sha256="f" * 64, commit=True
    )
    assert (out.imported, out.updated, out.unchanged) == (3, 2, 0)
    rows = await _customers(db, tenants.a)
    known = rows["971501234567"]
    assert known.name == "Ahmed Ali"  # blank → filled
    assert known.area == "Al Qusais"  # the customer's own correction is kept
    assert known.opt_in_status == "opted_in"
    assert rows["971561112233"].opt_in_status == "opted_out"  # an import never re-subscribes
    new = rows["971551234567"]
    assert (new.source, new.opt_in_status, new.opt_in_at) == ("excel_import", "pending", None)
    assert ("971501234567", "kept existing area") in out.conflicts

    # tenant B's customer with the same number is untouched
    assert (await _customers(db, tenants.b))["971551234567"].name == "B's customer"

    # the audit row holds counts and the file hash, nothing about a person
    async with db.tenant_session(tenants.a) as s:
        entry = await s.scalar(select(AuditLog).where(AuditLog.action == "import_customers"))
    assert entry is not None
    assert entry.after is not None
    assert entry.after["imported"] == 3
    assert "971501234567" not in str(entry.after)

    # second run: nothing changes
    before = {k: (v.name, v.area, v.emirate, v.address_note) for k, v in rows.items()}
    _, again = await import_excel.run(
        db, slug, table, MESSY_MAPPING, file_sha256="f" * 64, commit=True
    )
    assert (again.imported, again.updated, again.unchanged) == (0, 0, 5)
    after = await _customers(db, tenants.a)
    assert {k: (v.name, v.area, v.emirate, v.address_note) for k, v in after.items()} == before


def _big_workbook(path: Path, n: int, seed: int = 7) -> int:
    """n messy rows. Returns how many distinct valid mobiles it contains."""
    rng = random.Random(seed)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Aquarium Export"])
    ws.append(["S.No", "Name", "Mobile", "Area", "Emirate", "Address", "Joined"])
    areas = ["Al Nahda", "Al Qusais", "JLT", "Al Majaz", "Al Rashidiya", "النهدة", "القصيص"]
    first = ["Ahmed", "Mohammed", "Fatima", "Sara", "Omar", "Aisha", "John", "Priya",
             "محمد", "فاطمة", "أحمد", "خالد"]  # fmt: skip
    valid: set[str] = set()
    for i in range(n):
        base = 500000000 + rng.randrange(0, 8_000_000) * 7 % 90_000_000
        prefix = rng.choice(["50", "52", "54", "55", "56", "58"])
        local = prefix + str(base)[-7:]
        style = i % 9
        phone: Any
        if style == 0:
            phone = float(local)
        elif style == 1:
            phone = "0" + local
        elif style == 2:
            phone = f"+971 {local[:2]} {local[2:5]} {local[5:]}"
        elif style == 3:
            phone = "00971" + local
        elif style == 4:
            phone = ("0" + local).translate(str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩"))
        elif style == 5:
            phone = f"0{local[:2]}-{local[2:5]}-{local[5:]}"
        elif style == 6:
            phone = "04 " + local[2:]  # landline: rejected
        elif style == 7:
            phone = None if i % 2 else "n/a"  # rejected
        else:
            phone = int("971" + local)
        wa = normalize_phone(phone).wa_id
        if wa:
            valid.add(wa)
        if i % 500 == 250:
            ws.append([None] * 7)  # blank rows sprinkled through
        ws.append(
            [
                i + 1,
                f"{rng.choice(first)} {rng.choice(first)}".upper() if i % 3 == 0 else
                f"{rng.choice(first)} {rng.choice(first)}",
                phone,
                rng.choice(areas),
                rng.choice(["Dubai", "dubai", "DXB", "دبي", "Sharjah", "SHJ"]),
                f"Villa {rng.randrange(1, 400)}",
                datetime(2025, 1 + i % 12, 1 + i % 28),
            ]
        )  # fmt: skip
    wb.save(path)
    return len(valid)


async def test_ten_thousand_rows_clean_report_and_rerun_changes_nothing(
    db: Database, tmp_path: Path
) -> None:
    tenant = await make_tenant(db, "bulk")
    slug = await _slug(db, tenant)
    path = tmp_path / "big.xlsx"
    distinct = _big_workbook(path, 10_000)
    report = tmp_path / "report.csv"
    mapping = tmp_path / "mapping.json"

    # first run through the CLI: no terminal → the suggested mapping, saved for next time
    code = await _cli(
        [str(path), "--slug", slug, "--commit", "--report", str(report),
         "--save-mapping", str(mapping)]
    )  # fmt: skip
    assert code == 0
    saved = Mapping.model_validate_json(mapping.read_text())
    assert saved.columns["Mobile"] == "phone"
    assert saved.columns["Joined"] == "ignore"
    async with db.tenant_session(tenant) as s:
        count = await s.scalar(select(func.count()).select_from(Customer))
        pending = await s.scalar(select(func.count()).where(Customer.opt_in_status == "pending"))
    assert count == distinct
    assert pending == distinct

    with report.open(encoding="utf-8-sig") as f:
        lines = list(csv.DictReader(f))
    rejected = [r for r in lines if r["kind"] == "rejected"]
    folded = sum(len(r["rows"].split()) - 1 for r in lines if r["kind"] == "merged_in_file")
    assert len(rejected) + distinct + folded == 10_000  # every row is accounted for
    assert {r["detail"] for r in rejected} == {
        "landline number (not on WhatsApp)",
        "no phone number",
    }

    # re-run with the saved mapping: zero changes
    snapshot = {
        k: (v.name, v.area, v.emirate, v.address_note)
        for k, v in (await _customers(db, tenant)).items()
    }
    code = await _cli(
        [str(path), "--slug", slug, "--commit", "--mapping", str(mapping), "--report", str(report)]
    )
    assert code == 0
    again = await _customers(db, tenant)
    assert {k: (v.name, v.area, v.emirate, v.address_note) for k, v in again.items()} == snapshot


async def _cli(argv: list[str]) -> int:
    """import_excel.main runs its own event loop; run it off this one."""
    import asyncio

    return await asyncio.to_thread(import_excel.main, argv)


async def test_cli_reports_errors_without_a_traceback(
    db: Database, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    path = tmp_path / "x.csv"
    path.write_text("Name,Mobile\nA,0501234567\n")
    assert await _cli([str(path), "--slug", "no-such-tenant"]) == 2
    assert "no tenant with slug" in capsys.readouterr().err
